"""Document ingestion: PDF / DOCX / PPTX / TXT -> cited chunks.

Every chunk keeps the page (PDF), slide (PPTX) or section (DOCX/TXT) number it
came from. That metadata is what makes grounded citations possible downstream,
so it is attached at extraction time and never dropped.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import config


@dataclass
class Chunk:
    text: str
    source_page: int          # page / slide / section index (1-based)
    source_label: str         # e.g. "p.41", "slide 3", "section 2"
    doc_name: str
    index: int = 0
    extra: dict = field(default_factory=dict)

    def citation(self) -> str:
        return f"[{self.doc_name} {self.source_label}]"


# --------------------------------------------------------------------------- #
# Per-format extractors -> list[(unit_number, unit_label, raw_text)]
# --------------------------------------------------------------------------- #
def _extract_pdf(path: Path):
    from pypdf import PdfReader

    out = []
    try:
        reader = PdfReader(str(path))
        for i, page in enumerate(reader.pages, start=1):
            try:
                raw = page.extract_text() or ""
            except Exception as pe:
                print(f"[PDF WARN] Failed to extract page {i} of {path.name}: {pe}")
                raw = ""
            raw = raw.replace("\x00", "").strip()
            out.append((i, f"p.{i}", raw))
    except Exception as e:
        print(f"[PDF ERROR] Failed reading {path.name}: {e}")
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(str(path))
            for i, page in enumerate(pdf, start=1):
                raw = page.get_textpage().get_text_range() or ""
                out.append((i, f"p.{i}", raw.replace("\x00", "").strip()))
        except Exception:
            pass
    return out


def _extract_docx(path: Path):
    import docx

    doc = docx.Document(str(path))
    sections, cur, sec_no = [], [], 1
    for para in doc.paragraphs:
        txt = para.text.strip()
        if not txt:
            continue
        is_heading = (para.style.name or "").lower().startswith("heading")
        if is_heading and cur:
            sections.append((sec_no, f"section {sec_no}", "\n".join(cur)))
            sec_no += 1
            cur = []
        cur.append(txt)
    if cur:
        sections.append((sec_no, f"section {sec_no}", "\n".join(cur)))
    return sections or [(1, "section 1", "")]


def _extract_pptx(path: Path):
    from pptx import Presentation

    prs = Presentation(str(path))
    out = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
        out.append((i, f"slide {i}", "\n".join(parts)))
    return out


_HEADING_RE = re.compile(
    r"^\s*(section|chapter|unit|lesson|part)\s+([0-9ivxlcIVXLC]+)\b[.:) -]*\s*(.*)$",
    re.IGNORECASE,
)


def _extract_txt(path: Path):
    text = path.read_text(encoding="utf-8", errors="ignore")
    blocks = [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]
    if not blocks:
        return [(1, "section 1", text)]
    out = []
    for i, b in enumerate(blocks, start=1):
        first = b.splitlines()[0] if b.splitlines() else ""
        m = _HEADING_RE.match(first)
        label = f"{m.group(1).lower()} {m.group(2)}" if m else f"block {i}"
        out.append((i, label, b))
    return out


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".txt": _extract_txt,
    ".md": _extract_txt,
}

SUPPORTED_EXTS = tuple(_EXTRACTORS)


# --------------------------------------------------------------------------- #
# Chunking
# --------------------------------------------------------------------------- #
def _clean(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_units(text: str) -> list[str]:
    """Prefer paragraph / heading boundaries over mid-sentence cuts."""
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(paras) <= 1:
        paras = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text) if s.strip()]
    return paras


def _pack(units: list[str], lo: int, hi: int) -> list[str]:
    chunks, buf, count = [], [], 0
    for u in units:
        w = len(u.split())
        if count and count + w > hi:
            chunks.append(" ".join(buf))
            buf, count = [], 0
        buf.append(u)
        count += w
        if count >= lo:
            chunks.append(" ".join(buf))
            buf, count = [], 0
    if buf:
        tail = " ".join(buf)
        if chunks and len(tail.split()) < lo // 2:
            chunks[-1] = chunks[-1] + " " + tail
        else:
            chunks.append(tail)
    return chunks


def ingest(path: str | Path) -> list[Chunk]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext not in _EXTRACTORS:
        raise ValueError(
            f"Unsupported file type {ext!r}. Supported: {', '.join(SUPPORTED_EXTS)}"
        )
    units = _EXTRACTORS[ext](path)
    chunks: list[Chunk] = []
    lo, hi = config.CHUNK_MIN_WORDS, config.CHUNK_MAX_WORDS
    for unit_no, unit_label, raw in units:
        raw = _clean(raw)
        if not raw:
            continue
        for piece in _pack(_split_units(raw), lo, hi):
            piece = piece.strip()
            if len(piece.split()) < 8:
                continue
            chunks.append(
                Chunk(
                    text=piece,
                    source_page=unit_no,
                    source_label=unit_label,
                    doc_name=path.name,
                    index=len(chunks),
                )
            )
    if not chunks:
        raise ValueError(f"No extractable text found in {path.name}.")
    return chunks


def ingest_many(paths) -> list[Chunk]:
    out: list[Chunk] = []
    for p in paths:
        out.extend(ingest(p))
    for i, c in enumerate(out):
        c.index = i
    return out


def extract_doc_metadata(path: str | Path) -> dict[str, Any]:
    """Inspect and extract document metadata immediately upon upload.
    Returns:
        filename, page_count, word_count, chunk_count, suggested_topic, preview_snippet, is_valid, error
    """
    p = Path(path)
    ext = p.suffix.lower()
    if ext not in _EXTRACTORS:
        return {
            "filename": p.name,
            "is_valid": False,
            "error": f"Unsupported format '{ext}'. Supported: {', '.join(SUPPORTED_EXTS)}",
            "page_count": 0,
            "word_count": 0,
            "chunk_count": 0,
            "suggested_topic": "",
            "preview_snippet": "",
        }

    try:
        units = _EXTRACTORS[ext](p)
        total_pages = len(units)
        combined_text = " ".join(u[2] for u in units).strip()
        words = combined_text.split()
        word_count = len(words)

        chunks = ingest(p)
        chunk_count = len(chunks)

        # Inferred topic heuristics
        suggested_topic = ""
        if units and units[0][2]:
            first_lines = [line.strip() for line in units[0][2].splitlines() if len(line.strip()) > 3]
            for line in first_lines[:5]:
                clean_line = re.sub(r"^[\#\/\*\=\-\s]+", "", line).strip()
                if 5 < len(clean_line) < 70 and not re.search(r"page\s+\d|copyright|all rights reserved|isbn", clean_line, re.I):
                    suggested_topic = clean_line
                    break

        if not suggested_topic:
            raw_stem = p.stem.replace("_", " ").replace("-", " ")
            raw_stem = re.sub(r"^[0-9a-f]{8}\s*", "", raw_stem)
            suggested_topic = raw_stem.title()

        preview = combined_text[:200] + "..." if len(combined_text) > 200 else combined_text

        return {
            "filename": p.name,
            "is_valid": True,
            "page_count": total_pages,
            "word_count": word_count,
            "chunk_count": chunk_count,
            "suggested_topic": suggested_topic,
            "preview_snippet": preview,
            "error": None,
        }
    except Exception as exc:
        return {
            "filename": p.name,
            "is_valid": False,
            "error": str(exc),
            "page_count": 0,
            "word_count": 0,
            "chunk_count": 0,
            "suggested_topic": p.stem.replace("_", " ").title(),
            "preview_snippet": "",
        }
